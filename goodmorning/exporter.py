"""
엑셀 내보내기 (기존 naver_news_crawler10.py와 같은 규칙)

- 모든 기사: 온라인 1행 (날짜 = 온라인 게재일, URL = 기사 URL)
- 영상 기사: 방송 1행 추가 (날짜·URL은 온라인과 동일)
- 지면 기사: 지면 1행 추가 (날짜 = 온라인 게재일 + 1일, URL = '-')
- 정렬: 날짜 내림차순, 같은 날짜는 지면 → 방송 → 온라인
"""

from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

COLUMNS = ["연번", "날짜", "언론사명", "구분", "제목", "URL"]
WIDTHS = [8, 15, 20, 12, 70, 70]


def to_rows(articles):
    rows = []
    for a in articles:
        online_date = datetime.fromisoformat(a["published_at"])
        date_text = online_date.strftime("%Y.%m.%d.")
        press = a.get("press") or ""
        url = a.get("url") or a.get("naver_link") or a.get("original_link")

        if a.get("print_page") and a["print_page"] != "온라인":
            rows.append({
                "날짜": (online_date + timedelta(days=1)).strftime("%Y.%m.%d."),
                "언론사명": press, "구분": a["print_page"], "제목": a["title"], "URL": "-",
            })
        if a.get("has_video"):
            rows.append({"날짜": date_text, "언론사명": press, "구분": "방송", "제목": a["title"], "URL": url})
        rows.append({"날짜": date_text, "언론사명": press, "구분": "온라인", "제목": a["title"], "URL": url})

    # 중복 제거: 온라인·방송은 (URL, 구분), 지면은 (날짜, 언론사명, 구분, 제목)
    seen = set()
    unique = []
    for r in rows:
        if r["구분"] in ("온라인", "방송"):
            key = ("d", r["URL"], r["구분"])
        else:
            key = ("p", r["날짜"], r["언론사명"], r["구분"], r["제목"])
        if key not in seen:
            seen.add(key)
            unique.append(r)

    def sort_key(r):
        type_order = 1 if r["구분"] == "방송" else 2 if r["구분"] == "온라인" else 0
        return (-int(r["날짜"].replace(".", "")), type_order, r["언론사명"], r["제목"])

    unique.sort(key=sort_key)
    for index, r in enumerate(unique, start=1):
        r["연번"] = index
    return unique


def build_workbook(articles):
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    ws.append(COLUMNS)

    for r in to_rows(articles):
        ws.append([r[c] for c in COLUMNS])

    for letter, width in zip("ABCDEF", WIDTHS):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row[:4]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[4:]:
            cell.alignment = Alignment(vertical="center")
        url_cell = row[5]
        if isinstance(url_cell.value, str) and url_cell.value.startswith(("http://", "https://")):
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
